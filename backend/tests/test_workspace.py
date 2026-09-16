import io
import json
import time
from unittest.mock import Mock
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pypdf import PdfReader
from pptx import Presentation
from app.models.artifacts import ArtifactRequest, ArtifactDraft
from app.services import documents, jobs, store, gemini, chunks, renderers
from app.services.auth import User, current_user
from app.core.config import settings


def draft_for(source):
    return {'title':'AI 治理研究','summary':'原始研究摘要。','sections':[{'heading':'核心分析','body':'2026 年有 20 件需審查。','citations':[{'document_id':source['id'],'block_id':'p2'}]}], 'limitations':'合成資料，不能代表真實政策。'}


def artifact_request(source,kind='pdf'):
    return ArtifactRequest(kind=kind,scope='answer',message_ids=['a1'],source_ids=[source['id']],language='zh-TW')


def test_document_preserves_original_and_cleaned_draft():
    result=documents.create(b'Original source content.','source.txt','text/plain',{'language':'en'},'Human corrected text.')
    doc=documents.get(result['id'])
    assert store.get_bytes(doc.original_key)==b'Original source content.'
    assert doc.blocks[0]['text']=='Original source content.'
    assert doc.draft[0]['content']=='Human corrected text.'
    assert doc.published_version==0 and doc.index_status=='unpublished'


def test_private_document_and_job_isolation(source):
    job=jobs.create(artifact_request(source))
    token=current_user.set(User('other','other@example.test'))
    try:
        with pytest.raises(HTTPException) as denied:documents.get(source['id'])
        assert denied.value.status_code==404
        with pytest.raises(HTTPException):jobs.get(job['id'])
        with pytest.raises(HTTPException):jobs.create(artifact_request(source))
    finally:current_user.reset(token)


def test_pdf_draft_confirmation_and_download_bytes(source,monkeypatch):
    monkeypatch.setattr(gemini,'generate',Mock(return_value=json.dumps(draft_for(source))))
    task=jobs.create(artifact_request(source))
    jobs.run(task['id'])
    job=jobs.get(task['id'])
    assert job.status=='awaiting_review' and job.result is None
    jobs.mutate(job.id,'confirm',job.revision,job.draft)
    jobs.run(job.id)
    result=jobs.get(job.id)
    assert result.status=='completed',result.error
    data=store.get_bytes(result.result['files'][0]['key'])
    assert len(PdfReader(io.BytesIO(data)).pages)>=1
    assert result.draft['sections'][0]['citations'][0]['document_id']==source['id']


def test_cancel_retry_expiry_and_delete(source):
    task=jobs.create(artifact_request(source))
    cancelled=jobs.mutate(task['id'],'cancel',task['revision'])
    jobs.run(task['id'])
    assert jobs.get(task['id']).status=='cancelled'
    with pytest.raises(HTTPException):jobs.mutate(task['id'],'retry',task['revision'])
    retried=jobs.mutate(task['id'],'retry',cancelled['revision'])
    assert retried['status']=='queued'
    with store.session() as db:
        db.get(store.Job,task['id']).expires_at=time.time()-1
        db.commit()
    with pytest.raises(HTTPException):jobs.get(task['id'])
    jobs.purge_expired()
    with store.session() as db:assert db.get(store.Job,task['id']) is None


def test_rate_limit(source,monkeypatch):
    monkeypatch.setattr(settings,'JOBS_PER_USER_HOUR',1)
    jobs.create(artifact_request(source))
    with pytest.raises(HTTPException) as exc:jobs.create(artifact_request(source))
    assert exc.value.status_code==429


def test_translation_completeness_and_number_preservation(source,monkeypatch):
    def translate(content,*args,**kwargs):
        block=json.loads(content)['blocks'][0]
        return json.dumps({'blocks':[{'id':block['id'],'text':block['text'],'cells':block['cells']}]})
    monkeypatch.setattr(gemini,'generate',translate)
    request=ArtifactRequest(kind='translation',scope='document',source_ids=[source['id']],language='ja')
    task=jobs.create(request);jobs.run(task['id']);job=jobs.get(task['id'])
    assert job.status=='completed',job.error
    assert len(job.draft['blocks'])==3 and job.draft['complete']
    assert jobs.create(request)['id']==job.id
    assert documents.get(source['id']).blocks[1]['text'].startswith('In 2026')


def test_partial_translation_never_completed(source,monkeypatch):
    def translate(content,*args,**kwargs):
        block=json.loads(content)['blocks'][0]
        return json.dumps({'blocks':[{'id':block['id'],'text':block['text'].replace('20 cases','25 cases')}]})
    monkeypatch.setattr(gemini,'generate',translate)
    task=jobs.create(ArtifactRequest(kind='translation',scope='document',source_ids=[source['id']],language='ja'))
    jobs.run(task['id']);job=jobs.get(task['id'])
    assert job.status=='failed' and job.result is None
    assert len(job.draft['blocks'])==1 and not job.draft['complete']


def test_source_version_change_rejects_cached_generation(source):
    task=jobs.create(artifact_request(source))
    doc=documents.get(source['id'])
    store.put_bytes(doc.original_key,b'Changed original version.',doc.mime)
    with pytest.raises(HTTPException) as exc:jobs.inputs(artifact_request(source),User('editor','editor@example.test'))
    assert exc.value.status_code==409


def test_chunk_save_conflict_merge_and_no_publish(source):
    doc=documents.get(source['id'])
    merged=[dict(id='merged',order=0,content='\n'.join(c['content'] for c in doc.draft),refs=[r for c in doc.draft for r in c['refs']])]
    saved=chunks.save(doc.id,doc.draft_revision,merged)
    assert len(saved['draft'])==1 and saved['published_version']==0
    with pytest.raises(HTTPException):chunks.save(doc.id,doc.draft_revision,merged)
    with pytest.raises(HTTPException) as exc:chunks.publish(doc.id,saved['draft_revision'])
    assert exc.value.status_code==503
    reset=chunks.save(doc.id,saved['draft_revision'],[],True)
    assert len(reset['draft'])==1
    assert {r['block_id'] for r in reset['draft'][0]['refs']}=={'p1','p2','p3'}
    assert all(b['text'] in reset['draft'][0]['content'] for b in doc.blocks)


def test_chunk_invalid_ranges_rejected(source):
    doc=documents.get(source['id']);bad=doc.draft
    bad[0]['refs'][0]['end']=99999
    with pytest.raises(ValueError):chunks.validate(bad,doc.blocks)


def test_statistical_chart_must_trace_actual_number(source):
    doc,blocks=documents.evidence(source['id'])
    value=draft_for(source)
    value['chart']={'type':'bar','title':'Cases','points':[{'label':'Cases','value':20,'unit':'cases','period':'2026','source':{'document_id':doc.id,'block_id':'p2'},'quote':blocks[1]['text']}]}
    jobs.validate_draft(ArtifactDraft.model_validate(value),blocks)
    value['chart']['points'][0]['value']=42
    with pytest.raises(ValueError):jobs.validate_draft(ArtifactDraft.model_validate(value),blocks)


def test_editable_pptx_and_multilingual_pdf(source):
    draft=draft_for(source)
    sources=[dict(id=source['id'],title=source['title'],version=source['original_hash'],language='en')]
    for language in ['zh-TW','en','ja']:
        draft['title']='City governance' if language=='en' else '城市治理 / 都市ガバナンス'
        data=renderers.pdf(draft,language,sources)
        assert PdfReader(io.BytesIO(data)).pages
        slides=Presentation(io.BytesIO(renderers.pptx(draft,language,sources)))
        assert any(shape.has_text_frame and draft['title'] in shape.text for slide in slides.slides for shape in slide.shapes)


def test_auth_rejects_anonymous_and_untrusted_origin(monkeypatch):
    from app.main import app
    with TestClient(app) as client:
        assert client.get('/api/auth/config').status_code==200
        assert client.get('/api/library').status_code==401
        assert client.post('/api/artifacts',json={}).status_code==401
        assert client.post('/api/auth/login',json={'credential':'x'}).status_code==403
        assert client.get('/api/health').status_code==200


def test_api_cookie_session_and_mcp_tool_contract(source,monkeypatch):
    from app.services import auth
    from app.main import app
    from app.mcp.server import mcp
    import app.main as main_module
    monkeypatch.setattr(main_module,'session_manager',None)
    import asyncio
    monkeypatch.setattr(auth,'verify_google',lambda token:User('editor','editor@example.test'))
    with TestClient(app) as client:
        response=client.post('/api/auth/login',json={'credential':'test'},headers={'origin':'http://localhost:5173'})
        assert response.status_code==200
        assert client.get('/api/library').status_code==200
        assert client.get('/api/library/'+source['id']).json()['id']==source['id']
    tools=asyncio.run(mcp.list_tools())
    ask=next(t for t in tools if t.name=='ask_city_ai_governance_rag')
    assert 'response_language' in ask.inputSchema['properties'] and 'language' not in ask.inputSchema['properties']


def test_pdf_keeps_cjk_and_escapes_untrusted_markup():
    rendered=renderers.pdf_text('城市 <img src="https://example.test">\nReview & approval')
    assert '<img' not in rendered and '&lt;img' in rendered
    assert '<font name="CJK">城市 ' in rendered or '<font name="CJK">城市</font>' in rendered
    assert '<br/>' in rendered and 'Review &amp; approval' in rendered
    assert renderers.pdf_text('中文，句號。')=='<font name="CJK">中文，句號。</font>'
    assert renderers.pdf_text('한국어')=='<font name="CJK">한국어</font>'


def test_removed_document_language_rejected_before_storage(monkeypatch):
    writes=[]
    monkeypatch.setattr(store,'put_bytes',lambda *args:writes.append(args))
    with pytest.raises(HTTPException) as error:
        documents.create(b'Synthetic policy text.','unsupported.txt','text/plain',{'language':'ar'})
    assert error.value.status_code==422 and writes==[]


def test_docx_extraction_keeps_heading_table_and_paragraph_order():
    from docx import Document
    word=Document()
    word.add_heading('Policy 2026',level=1)
    table=word.add_table(rows=2,cols=2)
    for row,values in zip(table.rows,[['Period','Cases'],['2026','20']]):
        for cell,value in zip(row.cells,values):cell.text=value
    word.add_paragraph('Automatic approval is prohibited.')
    data=io.BytesIO();word.save(data)
    blocks,warnings=documents.extract(data.getvalue(),'policy.docx')
    assert [b['kind'] for b in blocks]==['heading','table','paragraph']
    assert blocks[1]['cells']==[['Period','Cases'],['2026','20']]
    assert blocks[2]['text']=='Automatic approval is prohibited.' and not warnings


def test_scanned_page_warning_cannot_be_bypassed_for_full_translation(monkeypatch):
    from reportlab.pdfgen import canvas
    data=io.BytesIO();pdf=canvas.Canvas(data)
    pdf.drawString(40,750,'Text page with enough extractable policy information for verification.')
    pdf.showPage();pdf.showPage();pdf.save()
    created=documents.create(data.getvalue(),'mixed-scan.pdf','application/pdf',{'language':'en'})
    assert any(w.startswith('OCR_REQUIRED: pages 2') for w in created['warnings'])
    request=ArtifactRequest(kind='translation',scope='document',source_ids=[created['id']],language='ja',
                            acknowledge_extraction_limits=True)
    with pytest.raises(HTTPException) as error:jobs.create(request)
    assert error.value.status_code==422
