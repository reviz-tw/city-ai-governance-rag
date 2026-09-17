import copy
import io
import json
import zipfile
import pytest
from pptx import Presentation
from pydantic import ValidationError
from app.models.artifacts import ArtifactDraft, ArtifactRequest, DeckOutline, Slide
from app.services import jobs, renderers, slide_authoring, gemini, store


def slide(ref, **extra):
    return dict(layout='briefing',title='Review remains a human responsibility',takeaway='Staff check privacy before approval.',
        points=[dict(label='Review',detail='Examine the source and privacy risks.'),dict(label='Approval',detail='The responsible person makes the decision.')],
        citations=[ref],notes='The source requires human review. This does not establish a particular review sequence.',**extra)


def deck(ref):
    first=slide(ref)
    second={**slide(ref),'title':'A proposed review sequence','layout':'process','reasoning':'recommendation'}
    third={**slide(ref),'title':'What the original evidence establishes','layout':'evidence',
           'quote':'No automatic approval is allowed.','points':[dict(label='Human decision',detail='The passage rules out automatic approval.')]}
    return dict(title='Governance review',summary='Human responsibility and privacy.',audience_goal='Understand the review requirement.',
                slides=[first,second,third],limitations='Synthetic evidence; no measured outcome is provided.',coverage_note='')


def test_native_layouts_are_editable_and_keep_notes_citations_and_exact_page_count(source):
    ref=dict(document_id=source['id'],block_id='p2')
    value=deck(ref)
    value['slides'].append({**slide(ref),'layout':'comparison','title':'Two review responsibilities','points':[],
        'columns':['Dimension','Requirement'],'rows':[['Decision','Human approval'],['Data','Privacy review']]})
    quote='In 2026, 20 cases required review and 12 cases were approved.'
    numeric=dict(type='bar',title='Review and approval counts',points=[dict(label=label,value=number,unit='cases',period='2026',quote=quote,source=ref)
        for label,number in [('Reviewed',20),('Approved',12)]])
    value['slides'].append({**slide(ref),'layout':'chart','title':'Counts describe different stages','points':[], 'chart':numeric})
    draft=ArtifactDraft.model_validate(value)
    jobs.validate_draft(draft,[dict(document_id=source['id'],id='p2',text=quote+' No automatic approval is allowed.')],'pptx')
    sources=[dict(id=source['id'],title='Synthetic source',version=source['original_hash'])]
    data=renderers.pptx(draft.model_dump(),'en',sources)
    result=Presentation(io.BytesIO(data))
    assert len(result.slides)==len(renderers.slide_plan(value,sources))==5
    assert any(shape.has_table for shape in result.slides[3].shapes)
    assert any(shape.has_chart for shape in result.slides[4].shapes)
    assert source['id'] in result.slides[0].notes_slide.notes_text_frame.text
    assert value['slides'][0]['notes'] in result.slides[0].notes_slide.notes_text_frame.text
    assert 'Limitations' in '\n'.join(shape.text for shape in result.slides[-1].shapes if shape.has_text_frame)
    assert all(0 <= shape.left and shape.left+shape.width <= result.slide_width and 0 <= shape.top and shape.top+shape.height <= result.slide_height
               for page in result.slides for shape in page.shapes)
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        assert any(name.endswith('.xlsx') for name in package.namelist())


def test_layout_contract_rejects_hidden_content_and_ragged_comparisons(source):
    ref=dict(document_id=source['id'],block_id='p2')
    with pytest.raises(ValidationError):
        Slide.model_validate({**slide(ref),'columns':['Hidden content']})
    with pytest.raises(ValidationError):
        Slide.model_validate({**slide(ref),'layout':'comparison','points':[],'columns':['A','B'],'rows':[['A'],['B','C']]})


def test_all_slide_citations_and_verbatim_quotes_are_checked(source):
    ref=dict(document_id=source['id'],block_id='p2')
    value=deck(ref)
    blocks=[dict(document_id=source['id'],id='p2',text='No automatic approval is allowed.')]
    value['slides'][1]['citations']=[dict(document_id='private-other-doc',block_id='p1')]
    with pytest.raises(ValueError,match='Citation'):jobs.validate_draft(ArtifactDraft.model_validate(value),blocks,'pptx')
    value['slides'][1]['citations']=[ref]
    value['slides'][2]['quote']='Automatic approval is allowed.'
    with pytest.raises(ValueError,match='verbatim'):jobs.validate_draft(ArtifactDraft.model_validate(value),blocks,'pptx')


def test_readable_layout_refuses_overflow_instead_of_truncating(source):
    ref=dict(document_id=source['id'],block_id='p2')
    value=deck(ref)
    value['slides'][0]['points']=[dict(label='完整保留內容',detail='不能默默刪除來源的重要限制條件。'*12) for _ in range(4)]
    value=ArtifactDraft.model_validate(value).model_dump()
    with pytest.raises(ValueError,match='readable slide area'):
        renderers.pptx(value,'zh-TW',[dict(id=source['id'],title='Source',version=source['original_hash'])])


def test_quotes_allow_pdf_line_wraps_but_preserve_words_and_qualifications(source):
    ref=dict(document_id=source['id'],block_id='p2')
    value=deck(ref)
    value['slides'][2]['quote']='宜假設任何輸入內容可能被利用。'
    blocks=[dict(document_id=source['id'],id='p2',text='宜假設任何輸\n入內容可能被利用。')]
    jobs.validate_draft(ArtifactDraft.model_validate(value),blocks,'pptx')
    value['slides'][2]['quote']='禁止輸入內容。'
    with pytest.raises(ValueError,match='verbatim'):
        jobs.validate_draft(ArtifactDraft.model_validate(value),blocks,'pptx')
    assert jobs.quotation_text('the rapist') != jobs.quotation_text('therapist')


def test_outline_requires_known_evidence_and_explains_a_shorter_deck(source):
    ref=dict(document_id=source['id'],block_id='p2')
    outline=DeckOutline.model_validate(dict(audience_goal='Understand review',slides=[dict(title='Review',purpose='Explain',layout='briefing',citations=[ref])]))
    blocks=[dict(document_id=source['id'],id='p2',text='Evidence')]
    with pytest.raises(ValueError,match='page count'):slide_authoring.validate_outline(outline,blocks,6)
    outline.coverage_note='Only one passage is available.'
    slide_authoring.validate_outline(outline,blocks,6)
    outline.slides[0].citations[0].block_id='unknown'
    with pytest.raises(ValueError,match='outside'):slide_authoring.validate_outline(outline,blocks,6)


def test_new_pptx_jobs_load_skill_plan_compose_review_and_preserve_edits(source,monkeypatch):
    ref=dict(document_id=source['id'],block_id='p2')
    value=deck(ref)
    stages=[]
    def generate(content,system,schema,**kwargs):
        assert 'Research slides' in system and 'Examples of evidence' in system
        stages.append(schema.__name__)
        if schema.__name__=='DeckOutline':
            return json.dumps(dict(audience_goal=value['audience_goal'],slides=[dict(title=s['title'],purpose='Explain the evidence',layout=s['layout'],citations=[ref]) for s in value['slides']]))
        if schema.__name__=='SlideReview':return json.dumps({'slides':[{'slide_number':i+1,'supported':True,'assessment':'Source supports the claim.'} for i in range(3)],'issues':[]})
        return json.dumps(value)
    monkeypatch.setattr(gemini,'generate',generate)
    request=ArtifactRequest(kind='pptx',scope='answer',message_ids=['a1'],source_ids=[source['id']],pages=3,language='en')
    task=jobs.create(request)
    assert jobs.get(task['id']).payload['slide_skill_version']==slide_authoring.version()
    jobs.run(task['id']);job=jobs.get(task['id'])
    assert job.status=='awaiting_review',job.error
    assert stages==['DeckOutline','SlideDeckDraft','SlideReview']
    job.draft['slides'][0]['title']='A reviewed title'
    monkeypatch.setattr(renderers,'slides_pdf',lambda data:b'%PDF-test-render')
    jobs.mutate(job.id,'confirm',job.revision,job.draft)
    jobs.run(job.id);result=jobs.get(job.id)
    assert result.status=='completed',result.error
    pptx=Presentation(io.BytesIO(store.get_bytes(result.result['files'][0]['key'])))
    assert any('A reviewed title' in shape.text for shape in pptx.slides[0].shapes if shape.has_text_frame)


def test_review_repairs_once_and_refuses_persistently_unsupported_content(source,monkeypatch):
    ref=dict(document_id=source['id'],block_id='p2')
    value=deck(ref)
    request=ArtifactRequest(kind='pptx',scope='answer',message_ids=['a1'],source_ids=[source['id']],pages=3)
    calls=[]
    def generate(content,system,schema,**kwargs):
        calls.append(schema.__name__)
        if schema.__name__=='DeckOutline':
            return json.dumps(dict(audience_goal='Understand',slides=[dict(title=s['title'],purpose='Explain',layout=s['layout'],citations=[ref]) for s in value['slides']]))
        if schema.__name__=='SlideReview':return json.dumps(dict(slides=[dict(slide_number=i+1,supported=False,assessment='Unsupported outcome.') for i in range(3)],issues=['Slide 1 invents an outcome.']))
        return json.dumps(value)
    monkeypatch.setattr(gemini,'generate',generate)
    with pytest.raises(ValueError,match='after one repair'):
        slide_authoring.generate(request,[dict(document_id=source['id'],id='p2',text='Evidence')],lambda draft:None,lambda p:None)
    assert calls==['DeckOutline','SlideDeckDraft','SlideReview','SlideDeckDraft','SlideReview']


def test_pdf_cannot_silently_accept_slide_only_draft(source):
    ref=dict(document_id=source['id'],block_id='p2')
    with pytest.raises(ValueError,match='per-slide'):
        jobs.validate_draft(ArtifactDraft.model_validate(deck(ref)),[], 'pdf')
