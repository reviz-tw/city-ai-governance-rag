import hashlib
import io
import re
import uuid
from pathlib import Path
from collections.abc import Mapping
from fastapi import HTTPException
from pypdf import PdfReader
from docx import Document as WordDocument
from sqlalchemy import select
from app.services import store
from app.services.auth import User, require_user, require_editor
from app.services.languages import normalize_language, detect


def can_read(document: store.Document, user: User):
    return user.admin or document.owner == user.id or document.shared or user.email in document.readers


def get(document_id: str, user: User | None = None, *, edit=False):
    user = user or require_user()
    with store.session() as db:
        doc = db.get(store.Document, document_id)
        if doc is None or not can_read(doc, user):
            raise HTTPException(404, 'Document not found')
        if edit and (not user.editor or not (user.admin or doc.owner == user.id)):
            raise HTTPException(403, 'Document editing not permitted')
        return doc


def extract(data: bytes, filename: str):
    blocks, warnings = [], []
    def add(text, page=None, kind='paragraph', cells=None):
        if text.strip():
            blocks.append(dict(id=f'p{len(blocks)+1}', text=text.strip(), page=page,
                               kind=kind, cells=cells))
    ext = Path(filename).suffix.lower()
    if ext == '.pdf':
        reader = PdfReader(io.BytesIO(data))
        empty = []
        for page_num, page in enumerate(reader.pages, 1):
            text = page.extract_text(extraction_mode='layout') or ''
            if len(re.sub(r'\s', '', text)) < 30:
                empty.append(page_num)
            for paragraph in re.split(r'\n\s*\n', text):
                add(paragraph, page_num)
        if empty:
            warnings.append('OCR_REQUIRED: pages ' + ', '.join(map(str, empty)))
        warnings.append('PDF_LAYOUT_REVIEW: extracted text may not preserve complex tables or images')
    elif ext == '.docx':
        doc = WordDocument(io.BytesIO(data))
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        for child in doc.element.body:
            if child.tag == qn('w:p'):
                paragraph = Paragraph(child, doc)
                add(paragraph.text, kind='heading' if paragraph.style.name.startswith('Heading') else 'paragraph')
            elif child.tag == qn('w:tbl'):
                rows = [[cell.text for cell in row.cells] for row in Table(child, doc).rows]
                add('\n'.join('\t'.join(row) for row in rows), kind='table', cells=rows)
    elif ext in {'.txt', '.md'}:
        for paragraph in re.split(r'\n\s*\n', data.decode('utf-8')):
            add(paragraph, kind='heading' if paragraph.startswith('#') else 'paragraph')
    else:
        raise HTTPException(415, 'Supported: PDF, DOCX, UTF-8 TXT and Markdown')
    if not blocks:
        raise HTTPException(422, 'No extractable text; OCR is required')
    if sum(len(b['text']) for b in blocks) > 500000:
        raise HTTPException(413, 'Document exceeds the text limit')
    return blocks, warnings


def baseline(blocks, limit=1500):
    chunks = []
    for block in blocks:
        text = block['text']
        start = 0
        while start < len(text):
            end = min(start + limit, len(text))
            if end < len(text):
                boundary = max(text.rfind('\n', start + limit // 2, end),
                               text.rfind('。', start + limit // 2, end),
                               text.rfind('. ', start + limit // 2, end))
                if boundary > start:
                    end = boundary + 1
            value = text[start:end]
            if value.strip():
                chunks.append(dict(id=f'c{len(chunks)+1}', order=len(chunks), content=value,
                                   refs=[dict(block_id=block['id'], start=start, end=end, page=block['page'])],
                                   algorithm='paragraph-v1'))
            start = end
    # Pack short adjacent paragraphs while retaining exact original source ranges.
    packed=[]
    for chunk in chunks:
        if packed and len(packed[-1]['content'])+2+len(chunk['content'])<=limit:
            packed[-1]['content']+='\n\n'+chunk['content']
            packed[-1]['refs'].extend(chunk['refs'])
        else:
            packed.append(dict(chunk,id=f'c{len(packed)+1}',order=len(packed),algorithm='paragraph-v2'))
    return packed


def create(data: bytes, filename: str, mime: str, metadata: dict, cleaned_text: str | None = None):
    user = require_editor()
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(413, 'Upload limit: 20 MiB')
    blocks, warnings = extract(data, filename)
    try:
        language = normalize_language(metadata.get('language') or detect('\n'.join(b['text'] for b in blocks)[:4000]) or 'zh-TW')
    except ValueError:
        raise HTTPException(422, 'Unsupported document language') from None
    document_id = uuid.uuid4().hex
    key = f'managed-originals/{document_id}/{hashlib.sha256(data).hexdigest()}/{Path(filename).name}'
    store.put_bytes(key, data, mime)
    draft = baseline(blocks)
    if cleaned_text is not None and cleaned_text.strip():
        # Human cleaned text is a draft; the extraction remains the immutable source.
        draft = baseline([dict(id='cleaned',text=cleaned_text,page=None)])
        for chunk in draft:
            chunk.update(refs=[dict(block_id=b['id'],start=0,end=len(b['text']),page=b['page']) for b in blocks],
                         algorithm='human-cleaned-paragraph-v1')
    with store.session() as db:
        doc = store.Document(id=document_id, owner=user.id, shared=False, readers=[], title=filename,
                             language=language, original_key=key, original_hash=hashlib.sha256(data).hexdigest(),
                             mime=mime, metadata_json=metadata, blocks=blocks, extraction_warnings=warnings,
                             draft=draft)
        db.add(doc)
        db.commit()
    return describe(doc)


def describe(doc, include_content=False):
    user = require_user()
    editable = user.editor and (user.admin or doc.owner == user.id)
    result = dict(id=doc.id, title=doc.title, language=doc.language, original_hash=doc.original_hash,
                  metadata=doc.metadata_json, warnings=doc.extraction_warnings,
                  draft_revision=doc.draft_revision, published_version=doc.published_version,
                  index_status=doc.index_status, shared=doc.shared, editable=editable)
    if editable:
        from app.core.config import settings
        result.update(readers=doc.readers, indexing_enabled=settings.CHUNK_INDEX_ENABLED and bool(settings.CHUNK_DATA_STORE_ID))
    if include_content:
        result.update(blocks=doc.blocks)
        if editable:
            result['draft'] = doc.draft
    return result


def list_documents():
    user = require_user()
    with store.session() as db:
        return [describe(d) for d in db.scalars(select(store.Document)) if can_read(d, user)]


def register_legacy(result):
    """Read-through catalogue of the pre-existing publicly served GCS collection."""
    require_user()
    link = result.get('link', '')
    from app.core.config import settings
    prefix = f'gs://{settings.GCS_BUCKET_NAME}/documents/'
    if not link.startswith(prefix):
        return None
    identifier = 'legacy-' + hashlib.sha256(link.encode()).hexdigest()[:32]
    with store.session() as db:
        if db.get(store.Document, identifier):
            return identifier
        filename = link.removeprefix(prefix)
        # Only this historical public collection is imported as shared.
        data = store.get_bytes(link)
        blocks, warnings = extract(data, filename)
        meta = result.get('metadata', {})
        db.add(store.Document(id=identifier, owner='legacy-public-collection', shared=True,
                             title=result['title'], language=normalize_language(meta.get('language') or detect('\n'.join(b['text'] for b in blocks)[:4000]) or 'zh-TW'),
                             original_key=link, original_hash=hashlib.sha256(data).hexdigest(),
                             mime='application/pdf' if filename.lower().endswith('.pdf') else 'application/octet-stream',
                             metadata_json=meta, blocks=blocks, extraction_warnings=warnings,
                             draft=baseline(blocks), index_status='legacy-index', published_version=0))
        db.commit()
    return identifier


def evidence(document_id, user=None):
    doc = get(document_id, user)
    data = store.get_bytes(doc.original_key)
    current_hash = hashlib.sha256(data).hexdigest()
    if current_hash != doc.original_hash:
        raise HTTPException(409, 'Original file changed; re-import and review the new source version')
    # Re-read the authorized original representation, never the AI answer.
    return doc, [dict(**b, document_id=doc.id, version=doc.original_hash) for b in doc.blocks]
