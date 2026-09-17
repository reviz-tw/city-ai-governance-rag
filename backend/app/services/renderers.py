"""Fixed, escaped templates only. No generated code, remote URLs or HTML execution."""
import io
import json
import math
import textwrap
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from app.core.config import settings


def font_path(cjk=False):
    names = ([str(Path(settings.FONT_DIR) / 'cjk.ttc'), '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc',
              '/System/Library/Fonts/Supplemental/Arial Unicode.ttf'] if cjk else
             [str(Path(settings.FONT_DIR) / 'sans.ttf'), '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
              '/System/Library/Fonts/Supplemental/Arial.ttf'])
    return next((p for p in names if Path(p).exists()), None)


def register_fonts():
    for name, cjk in [('Body', False), ('CJK', True)]:
        if name in pdfmetrics.getRegisteredFontNames():
            continue
        path = font_path(cjk)
        if path:
            pdfmetrics.registerFont(TTFont(name, path))
        else:
            raise RuntimeError(f'Missing required {name} font; install deployed renderer fonts')


def cjk_character(char):
    return ('\u3000' <= char <= '\u9fff' or '\uac00' <= char <= '\ud7af' or
            '\uff00' <= char <= '\uffef')


def pdf_text(text):
    # Preserve embedded CJK text while escaping every user-provided text run.
    result = []
    for line in text.split('\n'):
        pieces, current, previous = [], '', None
        for char in line:
            font = 'CJK' if cjk_character(char) else 'Body'
            if previous and font != previous:
                pieces.append(f'<font name="{previous}">{escape(current)}</font>')
                current = ''
            current += char
            previous = font
        if current:
            pieces.append(f'<font name="{previous}">{escape(current)}</font>')
        result.append(''.join(pieces))
    return '<br/>'.join(result)


def references(sources):
    return '\n'.join(f'[{i+1}] {s["title"]} | {s["id"]} | {s["version"]} | {s["language"]}'
                     for i, s in enumerate(sources))


def pdf(draft, language, sources):
    register_fonts()
    output = io.BytesIO()
    body = ParagraphStyle('body', fontName='Body', fontSize=10, leading=17, spaceAfter=9,
                          wordWrap='CJK', alignment=0)
    heading = ParagraphStyle('heading', parent=body, fontSize=16, leading=23, spaceBefore=14,
                             textColor=colors.HexColor('#164E63'), keepWithNext=True)
    title = ParagraphStyle('title', parent=heading, fontSize=23, leading=31)
    flow = [Paragraph(pdf_text(draft['title']), title),
            Paragraph(pdf_text(draft['summary']), body)]
    for section in draft['sections']:
        flow.extend([Paragraph(pdf_text(section['heading']), heading),
                     Paragraph(pdf_text(section['body']), body)])
        if section.get('citations'):
            flow.append(Paragraph(pdf_text(' | '.join(f'{c["document_id"]}/{c["block_id"]}' for c in section['citations'])), body))
    flow.extend([Paragraph(pdf_text(draft['limitations']), body),
                 Paragraph(pdf_text(references(sources)), body),
                 Paragraph(pdf_text(datetime.now(timezone.utc).isoformat()), body)])
    def footer(canvas, doc):
        canvas.setFont('Body', 8)
        canvas.drawRightString(550, 28, str(doc.page))
    SimpleDocTemplate(output, pagesize=(595.28, 841.89), rightMargin=42, leftMargin=42,
                      topMargin=42, bottomMargin=42, title=draft['title']).build(flow, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def slide_plan(draft, sources):
    if draft.get('slides'):
        return draft['slides']
    pages=[]
    def add(title,content,refs=''):
        limit=450 if any(cjk_character(char) for char in content) else 800
        parts=textwrap.wrap(content,width=limit,break_long_words=True,replace_whitespace=False) or ['']
        for index,part in enumerate(parts):
            pages.append((title+(f' ({index+1})' if len(parts)>1 else ''),part,refs))
    add(draft['title'],draft['summary'])
    for section in draft['sections']:
        refs=' | '.join(f'{c["document_id"]}/{c["block_id"]}' for c in section.get('citations',[]))
        add(section['heading'],section['body'],refs)
    add('Limitations / 限制',draft['limitations'])
    source_lines=references(sources).split('\n')
    for offset in range(0,len(source_lines),3):
        add('References / 參考來源','\n\n'.join(source_lines[offset:offset+3]))
    return pages


def pptx(draft, language, sources):
    if draft.get('slides'):
        from app.services.slide_rendering import render
        return render(draft, language, sources)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    def slide(title, content, refs=''):
        page = deck.slides.add_slide(deck.slide_layouts[6])
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = RGBColor.from_string('F7F8F6')
        for text, x, y, w, h, size, color in [
            (title, .7, .4, 11.9, 1.5, 20 if len(title)>75 else 28, '164E63'),
            (content, .7, 2.0, 11.9, 4.3, 22 if len(content) < 650 else 16, '20313A'),
            (refs, .7, 6.65, 11.9, .55, 9, '52636B')]:
            box = page.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            box.text_frame.word_wrap = True
            for i, line in enumerate(text.split('\n')):
                paragraph = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
                paragraph.text = line
                cjk_font = 'WenQuanYi Zen Hei' if '/usr/share/' in (font_path(True) or '') else 'Arial Unicode MS'
                sans_font = 'DejaVu Sans' if '/usr/share/' in (font_path(False) or '') else 'Arial'
                paragraph.font.name = sans_font
                for run in paragraph.runs:
                    rpr = run._r.get_or_add_rPr()
                    for tag, face in [('a:ea', cjk_font), ('a:cs', sans_font)]:
                        font_element = OxmlElement(tag)
                        font_element.set('typeface',face)
                        rpr.append(font_element)
                paragraph.font.size = Pt(size)
                paragraph.font.color.rgb = RGBColor.from_string(color)
                paragraph.space_after = Pt(9)
        page.notes_slide.notes_text_frame.text = refs
    for title,content,refs in slide_plan(draft,sources):
        slide(title,content,refs)
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def chart(spec, language, sources):
    """Fixed SVG/PNG primitives with measured wrapping and actual diagram connectors."""
    width=1200
    rows=spec.get('points') if spec['type']=='bar' else spec.get('rows') if spec['type']=='comparison' else spec.get('labels')
    if not rows:
        raise ValueError('Chart has no data')
    fonts={}
    def font(size,cjk):
        key=(size,cjk)
        if key not in fonts:
            path=font_path(cjk)
            if not path:
                raise RuntimeError('Missing chart font')
            fonts[key]=ImageFont.truetype(path,size)
        return fonts[key]
    def runs(text,size):
        values=[]
        for char in text:
            cjk=cjk_character(char)
            if values and values[-1][1]==cjk:
                values[-1]=(values[-1][0]+char,cjk)
            else:
                values.append((char,cjk))
        return [(text,font(size,cjk)) for text,cjk in values]
    def measure(text,size):
        return sum(face.getlength(value) for value,face in runs(text,size))
    def wrap(text,size,limit):
        lines=[]
        for paragraph in str(text).split('\n'):
            line=''
            for char in paragraph:
                if line and measure(line+char,size)>limit:
                    boundary=line.rfind(' ')
                    if boundary>len(line)//2:
                        lines.append(line[:boundary]);line=line[boundary+1:]+char
                    else:
                        lines.append(line);line=char
                else:
                    line+=char
            lines.append(line)
        return lines
    title_lines=wrap(spec['title'],30,1120)
    row_texts=[]
    for index,row in enumerate(rows):
        if spec['type']=='bar':
            text=f'{row["label"]}\n{row["value"]:g} {row["unit"]} · {row["period"]}'
        elif spec['type']=='comparison':
            text=' | '.join(row)
        else:
            text=f'{index+1}. {row}'
        row_texts.append(wrap(text,22,1010 if spec['type'] in {'flow','architecture'} else 1120))
    positions=[]
    y=45+len(title_lines)*38+30
    for lines in row_texts:
        height=max(80,len(lines)*30+30)+(40 if spec['type']=='bar' else 0)
        positions.append((y,height));y+=height+24
    footer=wrap(references(sources),13,1120)
    height=int(y+len(footer)*19+75)
    image=Image.new('RGB',(width,height),'#f7f8f6');draw=ImageDraw.Draw(image)
    svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><rect width="100%" height="100%" fill="#f7f8f6"/>']
    def label(lines,x,y,size=22,right=1160):
        for index,line in enumerate(lines):
            top=y+index*(size+8)
            left=x
            for value,face in runs(line,size):
                draw.text((left,top),value,font=face,fill='#164e63');left+=face.getlength(value)
            svg.append(f'<text x="{x}" y="{top+size}" font-family="DejaVu Sans,WenQuanYi Zen Hei,sans-serif" font-size="{size}" fill="#164e63">{escape(line)}</text>')
    def rectangle(x,y,w,h,color):
        draw.rectangle((x,y,x+w,y+h),fill=color)
        svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="5" fill="{color}"/>')
    label(title_lines,40,30,30)
    max_value=max((p['value'] for p in rows),default=1) if spec['type']=='bar' else 1
    for index,lines in enumerate(row_texts):
        top,h=positions[index]
        if spec['type']=='bar':
            label(lines,40,top)
            rectangle(40,top+len(lines)*30+12,1050*rows[index]['value']/(max_value or 1),25,'#168a8a')
        elif spec['type']=='comparison':
            label(lines,40,top)
        else:
            rectangle(35,top,1060,h,'#dceceb');label(lines,50,top+15,right=1075)
    if spec['type'] in {'flow','architecture'}:
        for start,end in spec.get('edges',[]):
            y1=positions[start][0]+positions[start][1]/2;y2=positions[end][0]+positions[end][1]/2
            points=[(1095,y1),(1130,y1),(1130,y2),(1095,y2)]
            draw.line(points,fill='#168a8a',width=3)
            draw.polygon([(1095,y2),(1106,y2-6),(1106,y2+6)],fill='#168a8a')
            svg.append(f'<polyline points="{" ".join(f"{x},{y}" for x,y in points)}" stroke="#168a8a" stroke-width="3" fill="none"/><polygon points="1095,{y2} 1106,{y2-6} 1106,{y2+6}" fill="#168a8a"/>')
    label(footer,40,y+15,13)
    label([datetime.now(timezone.utc).date().isoformat()],40,height-30,12)
    svg.append('</svg>');output=io.BytesIO();image.save(output,format='PNG')
    return ''.join(svg).encode(),output.getvalue()


def slides_pdf(data):
    import subprocess
    import tempfile
    with tempfile.TemporaryDirectory(prefix='governance-render-') as directory:
        root = Path(directory)
        source = root / 'slides.pptx'
        source.write_bytes(data)
        font_config = root / 'fonts.conf'
        dirs = sorted({str(Path(p).parent) for p in [font_path(False), font_path(True)] if p})
        font_config.write_text('<fontconfig>' + ''.join('<dir>'+escape(d)+'</dir>' for d in dirs) + '<cachedir>'+str(root / 'font-cache')+'</cachedir></fontconfig>')
        command = [settings.SOFFICE_BIN, f'-env:UserInstallation={(root / "profile").as_uri()}',
                   '--headless', '--convert-to', 'pdf', '--outdir', str(root), str(source)]
        result = subprocess.run(command, timeout=90, capture_output=True, check=False,
                                env={**__import__('os').environ, 'SAL_DISABLE_OPENCL':'1', 'FONTCONFIG_FILE':str(font_config)})
        output = root / 'slides.pdf'
        if result.returncode or not output.exists():
            raise RuntimeError('Slide PDF rendering failed')
        return output.read_bytes()
