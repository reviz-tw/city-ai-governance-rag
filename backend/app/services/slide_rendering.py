"""Editable, measured slide layouts. Content comes from the slide schema, never model code."""
import io
import re
from functools import lru_cache
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement
from app.services.renderers import font_path, cjk_character

INK, ACCENT, MUTED, PAPER, RULE = '20313A', '164E63', '52636B', 'F7F8F6', 'D3DDD9'
LABELS = {
    'zh': ('資料依據', '推論', '建議', '限制', '來源'),
    'zh-CN': ('资料依据', '推论', '建议', '限制', '来源'),
    'en': ('Evidence', 'Inference', 'Recommendation', 'Limitations', 'Sources'),
    'ja': ('根拠', '推論', '提案', '制約', '出典'),
    'fr': ('Constats', 'Interprétation', 'Recommandation', 'Limites', 'Sources'),
    'es': ('Evidencia', 'Inferencia', 'Recomendación', 'Limitaciones', 'Fuentes'),
    'ru': ('Данные', 'Вывод', 'Рекомендация', 'Ограничения', 'Источники'),
    'ko': ('근거', '추론', '권고', '한계', '출처'),
    'de': ('Belege', 'Schlussfolgerung', 'Empfehlung', 'Grenzen', 'Quellen'),
}


@lru_cache(maxsize=64)
def font(size, cjk):
    path = font_path(cjk)
    if not path:
        raise RuntimeError('Missing slide font')
    return ImageFont.truetype(path, size*4)


def wrap(text, width, size):
    """Explicit line breaks using the same installed font as the native text object."""
    cjk = any(cjk_character(c) for c in text)
    face = font(size, cjk)
    max_width = (width*72 - 8)*.95
    lines = []
    for paragraph in text.split('\n'):
        line = ''
        tokens = re.findall(r'[\u2e80-\u9fff\uac00-\ud7af][，。！？；：、）》」』】]*|[^\s\u2e80-\u9fff\uac00-\ud7af]+|\s+', paragraph)
        for token in tokens:
            if face.getlength(line+token)/4 <= max_width:
                line += token
                continue
            if line.strip():
                lines.append(line.rstrip())
            line = token.lstrip()
            if face.getlength(line)/4 > max_width:
                current = ''
                for char in line:
                    if face.getlength(current+char)/4 > max_width:
                        lines.append(current)
                        current = ''
                    current += char
                line = current
        lines.append(line.rstrip())
    if cjk and len(lines)==2 and '\n' not in text and face.getlength(lines[-1])/4 < max_width*.35:
        # Balance CJK text instead of leaving a character or punctuation alone.
        joined=''.join(lines)
        choices=[i for i in range(1,len(joined)) if joined[i] not in '，。！？；：、）》」』】' and not (joined[i-1].isascii() and joined[i-1].isalnum() and joined[i].isascii() and joined[i].isalnum())]
        if choices:
            middle=min(choices,key=lambda i:abs(face.getlength(joined[:i])-face.getlength(joined[i:])))
            lines=[joined[:middle],joined[middle:]]
    return lines, cjk


def format_text(frame, text, width, height, size=22, minimum=18, color=INK, bold=False):
    for chosen in range(size, minimum-1, -1):
        lines, cjk = wrap(text, width, chosen)
        if len(lines)*chosen*1.20 <= height*72-8:
            break
    else:
        raise ValueError(f'Text exceeds the readable slide area; shorten or move detail to notes: {text[:80]}')
    frame.clear()
    # Explicitly measured breaks prevent LibreOffice/Slides adding a second wrap.
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = Pt(4)
    frame.margin_top = frame.margin_bottom = Pt(4)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    name = ('WenQuanYi Zen Hei' if '/usr/share/' in (font_path(True) or '') else 'Arial Unicode MS') if cjk else ('DejaVu Sans' if '/usr/share/' in (font_path() or '') else 'Arial')
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = name
        paragraph.font.size = Pt(chosen)
        # Bundled CJK regular faces have no matching bold face; synthetic bold can
        # overlap glyphs in LibreOffice. Hierarchy comes from size and color.
        paragraph.font.bold = bold and not cjk
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = paragraph.space_before = Pt(0)
        paragraph.line_spacing = Pt(chosen*1.20)
        for run in paragraph.runs:
            run.font.name = name
            run.font.size = Pt(chosen)
            run.font.bold = bold and not cjk
            rpr = run._r.get_or_add_rPr()
            for tag in ('a:ea','a:cs'):
                node = OxmlElement(tag)
                node.set('typeface',name)
                rpr.append(node)


def textbox(page, text, x, y, w, h, size=22, minimum=18, color=INK, bold=False):
    shape = page.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    format_text(shape.text_frame,text,w,h,size,minimum,color,bold)
    return shape


def render(draft, language, sources):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = draft['title']
    deck.core_properties.subject = draft.get('audience_goal','')
    labels = LABELS.get(language,LABELS.get(language.split('-')[0],LABELS['en']))
    source_map = {source['id']:(index+1,source) for index,source in enumerate(sources)}
    for index, slide in enumerate(draft['slides']):
        page = deck.slides.add_slide(deck.slide_layouts[6])
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = RGBColor.from_string(PAPER)
        reasoning = labels[{'evidence':0,'inference':1,'recommendation':2}[slide['reasoning']]]
        textbox(page,reasoning,.65,.22,10,.38,14,14,MUTED)
        textbox(page,str(index+1),12,.22,.6,.38,14,14,MUTED)
        textbox(page,slide['title'],.65,.73,12,1.06,32,28,ACCENT,True)
        textbox(page,slide['takeaway'],.65,1.91,12,.83,23,20,INK)
        y, height = 2.82, 3.28
        layout = slide['layout']
        if layout == 'briefing':
            count = len(slide['points'])
            row_h = height/count
            for pos, point in enumerate(slide['points']):
                textbox(page,point['label'],.65,y+pos*row_h,3.2,row_h-.04,23,18,ACCENT,True)
                textbox(page,point['detail'],4.12,y+pos*row_h,8.5,row_h-.04,22,18)
        elif layout == 'process':
            count = len(slide['points'])
            row_h = height/count
            for pos, point in enumerate(slide['points']):
                yy = y+pos*row_h
                textbox(page,f'{pos+1:02}',.65,yy,.65,row_h-.06,24,20,ACCENT,True)
                textbox(page,point['label'],1.50,yy,3.1,row_h-.06,21,18,ACCENT,True)
                textbox(page,point['detail'],4.88,yy,7.7,row_h-.06,21,18)
                if pos < count-1:
                    line = page.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(.96),Inches(yy+row_h-.13),Inches(.96),Inches(yy+row_h+.01))
                    line.line.color.rgb = RGBColor.from_string(ACCENT)
                    line.line.width = Pt(1.5)
        elif layout == 'comparison':
            rows = [slide['columns'], *slide['rows']]
            row_h, col_w = height/len(rows), 12/len(slide['columns'])
            table = page.shapes.add_table(len(rows),len(slide['columns']),Inches(.65),Inches(y),Inches(12),Inches(height)).table
            for r,row in enumerate(rows):
                table.rows[r].height = Inches(row_h)
                for c,text in enumerate(row):
                    cell = table.cell(r,c)
                    cell.margin_left = cell.margin_right = Inches(.10)
                    cell.margin_top = cell.margin_bottom = Inches(.03)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(ACCENT if r==0 else ('FFFFFF' if r%2 else 'ECF0ED'))
                    format_text(cell.text_frame,text,col_w-.2,row_h-.06,20,18,'FFFFFF' if r==0 else INK,r==0)
        elif layout == 'evidence':
            textbox(page,slide['quote'],.65,y,6.05,height,26,20,ACCENT)
            row_h = height/len(slide['points'])
            for pos, point in enumerate(slide['points']):
                textbox(page,point['label'],7.08,y+pos*row_h,5.55,.57,23,19,ACCENT,True)
                textbox(page,point['detail'],7.08,y+pos*row_h+.69,5.55,row_h-.79,21,18)
        elif layout == 'chart':
            spec = slide['chart']
            data = CategoryChartData()
            data.categories = [point['label'] for point in spec['points']]
            data.add_series(spec['points'][0]['unit'],[point['value'] for point in spec['points']])
            chart = page.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.65),Inches(y),Inches(12),Inches(height),data).chart
            chart.has_legend = False
            chart.category_axis.tick_labels.font.size = Pt(18)
            chart.value_axis.tick_labels.font.size = Pt(16)
            chart_face = ('WenQuanYi Zen Hei' if '/usr/share/' in (font_path(True) or '') else 'Arial Unicode MS') if language.split('-')[0] in {'zh','ja','ko'} else ('DejaVu Sans' if '/usr/share/' in (font_path() or '') else 'Arial')
            chart.category_axis.tick_labels.font.name = chart_face
            chart.value_axis.tick_labels.font.name = chart_face
            chart.value_axis.minimum_scale = 0
            chart.value_axis.has_title = True
            chart.value_axis.axis_title.text_frame.text = f'{spec["points"][0]["unit"]} ({spec["points"][0]["period"]})'
            for paragraph in chart.value_axis.axis_title.text_frame.paragraphs:
                paragraph.font.name = chart_face
                paragraph.font.size = Pt(16)
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            plot.data_labels.font.size = Pt(18)
            plot.data_labels.font.name = chart_face
            plot.series[0].format.fill.solid()
            plot.series[0].format.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
            if any(len(wrap(point['label'],12/len(spec['points']),18)[0]) > 2 for point in spec['points']):
                raise ValueError('Shorten chart category labels; put explanations in notes')

        if index == len(draft['slides'])-1:
            textbox(page,f'{labels[3]}: {draft["limitations"]}',.65,6.10,12,.60,15,14,MUTED)
        visible_refs, full_refs = [], []
        seen_sources = set()
        for ref in slide['citations']:
            number, source = source_map[ref['document_id']]
            if number not in seen_sources:
                short_title = source['title'] if len(source['title']) <= 26 else source['title'][:25]+'…'
                visible_refs.append(f'[{number}] {short_title}')
                seen_sources.add(number)
            full_refs.append(f'[{number}] {source["title"]}\n{ref["document_id"]}/{ref["block_id"]}\nVersion: {source["version"]}')
        textbox(page,labels[4]+': '+', '.join(visible_refs),.65,6.89,12,.50,12,11,MUTED)
        page.notes_slide.notes_text_frame.text = (slide['notes']+'\n\n'+labels[4]+':\n'+'\n\n'.join(full_refs)
            +'\n\n'+labels[3]+': '+draft['limitations'])
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()
